"""Document parsers: bytes -> [(page number, text)]. Page numbers are 1-based; formats without pages use 1."""
from __future__ import annotations

import csv
import io
import re
from html.parser import HTMLParser

MAX_BYTES = 20 * 1024 * 1024
TYPES = {".pdf", ".docx", ".md", ".markdown", ".txt", ".csv", ".html", ".htm"}


def kind_of(name: str) -> str:
    m = re.search(r"\.[A-Za-z0-9]+$", name)
    ext = m.group(0).lower() if m else ""
    if ext not in TYPES:
        raise ValueError(f"unsupported file type {ext or '(none)'}: use PDF, DOCX, MD, TXT, CSV or HTML")
    return ext


def clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"-\n(?=[a-z])", "", text)        # words hyphenated across lines
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_pdf(data: bytes) -> list[tuple[int, str]]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return [(i + 1, clean(page.extract_text() or "")) for i, page in enumerate(reader.pages)]


def parse_docx(data: bytes) -> list[tuple[int, str]]:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return [(1, clean("\n\n".join(parts)))]


def parse_csv(data: bytes) -> list[tuple[int, str]]:
    rows = list(csv.reader(io.StringIO(data.decode("utf-8", "replace"))))
    if not rows:
        return [(1, "")]
    head, body = rows[0], rows[1:]
    lines = ["; ".join(f"{h}: {v}" for h, v in zip(head, r) if v) for r in body] if body else [", ".join(head)]
    return [(1, "\n".join(lines))]


class _Text(HTMLParser):
    SKIP = {"script", "style", "nav", "footer", "header", "noscript", "svg"}
    BLOCK = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "section", "article", "table"}

    def __init__(self):
        super().__init__()
        self.out: list[str] = []
        self.skip = 0
        self.title = ""
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        if tag in self.BLOCK:
            self.out.append("\n")
        self._in_title = tag == "title"

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        if tag in self.BLOCK:
            self.out.append("\n")
        self._in_title = False

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self.skip:
            self.out.append(data)


def parse_html(data: bytes) -> tuple[str, list[tuple[int, str]]]:
    p = _Text()
    p.feed(data.decode("utf-8", "replace"))
    return p.title.strip(), [(1, clean("".join(p.out)))]


def parse(name: str, data: bytes) -> list[tuple[int, str]]:
    if len(data) > MAX_BYTES:
        raise ValueError("file is larger than 20 MB")
    ext = kind_of(name)
    if ext == ".pdf":
        return parse_pdf(data)
    if ext == ".docx":
        return parse_docx(data)
    if ext == ".csv":
        return parse_csv(data)
    if ext in (".html", ".htm"):
        return parse_html(data)[1]
    return [(1, clean(data.decode("utf-8", "replace")))]


def chunk(pages: list[tuple[int, str]], size: int = 1800, overlap: int = 250) -> list[tuple[int, str]]:
    """~450-token chunks along paragraph boundaries, each tagged with the page it starts on."""
    out: list[tuple[int, str]] = []
    for page, text in pages:
        paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
        cur = ""
        for para in paras:
            while len(para) > size:                         # a giant paragraph: cut at sentence ends
                cut = para.rfind(". ", 0, size)
                cut = cut + 1 if cut > size // 2 else size
                piece, para = para[:cut].strip(), para[max(0, cut - overlap):].strip()
                if cur:
                    out.append((page, cur))
                    cur = ""
                out.append((page, piece))
            if cur and len(cur) + len(para) + 2 > size:
                out.append((page, cur))
                cur = cur[-overlap:].split(" ", 1)[-1] + "\n\n" + para if overlap else para
            else:
                cur = f"{cur}\n\n{para}" if cur else para
        if cur:
            out.append((page, cur))
    return out
